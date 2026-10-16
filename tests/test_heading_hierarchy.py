"""标题分级（M1.1）、标题栈（M1.2）与冻结字段（M1.3）的回归测试。

fixture 全部用 python-docx 现场生成，不依赖 ``data/`` 语料；正反例取自
改造方案 §7 ``tests/test_heading_hierarchy.py`` 一行。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from src.data_processing.doc_tree import node_id
from src.data_processing.parser_chunker import SyllabusChunker


class _FakeParagraph:
    """最小 stub：只提供 ``_heading_level`` 需要的 ``text`` / ``style.name``。"""

    def __init__(self, text: str, style_name: str = "Normal") -> None:
        self.text = text
        self.style = type("_Style", (), {"name": style_name})()


def _write_paragraphs(path: Path, paragraphs: list[str]) -> Path:
    """把一串段落写成一个临时 DOCX。"""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


def _fill_table(table: Any, rows: list[list[str]]) -> None:
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value


def test_heading_level_recognises_headings_and_rejects_body_lines(
    tmp_path: Path,
) -> None:
    positives = [
        ("一、课程基本信息", 1),
        ("四、实习内容与要求", 1),
        ("第3章 关系数据库模型", 1),
        ("（三）教学方法", 2),
        ("第3节 关系模型", 2),
        ("3.2 索引结构", 2),
        ("1. 教材", 3),
        ("2、话题讨论", 3),
    ]
    negatives = [
        "2. 平时成绩包括4次作业和考勤：……",
        "5. 邱冠华主编. 图书馆阅读推广基础工作. 北京：朝华出版社，2015",
        "3. 期末考核方式为闭卷考试。",
        "2.课程组织：对学生进行分组，并设立负责人。根据课程内容与课后作业，……",
        "一、总成绩构成：平时考核占比40%，期末考试占比60%。",
        "二、平时考核说明：",
        "这是很长的一段正文说明文字，不应该被识别成标题。",
    ]
    path = _write_paragraphs(
        tmp_path / "headings.docx",
        [text for text, _ in positives] + negatives,
    )
    paragraphs = Document(path).paragraphs
    chunker = SyllabusChunker()

    for paragraph, (text, expected) in zip(paragraphs, positives):
        assert chunker._heading_level(paragraph) == expected, text

    for paragraph in paragraphs[len(positives) :]:
        assert chunker._heading_level(paragraph) is None, paragraph.text

    assert chunker._is_heading(paragraphs[0]) is True
    assert chunker._is_heading(paragraphs[-1]) is False


def test_heading_level_uses_word_style_and_week_pattern(tmp_path: Path) -> None:
    path = tmp_path / "styled.docx"
    document = Document()
    document.add_paragraph("关系数据库", style="Heading 2")
    document.add_paragraph("查询优化", style="Heading 3")
    document.add_paragraph("第一周 课程导论")
    document.save(path)

    chunker = SyllabusChunker()
    paragraphs = Document(path).paragraphs
    assert chunker._heading_level(paragraphs[0]) == 2
    assert chunker._heading_level(paragraphs[1]) == 3
    assert chunker._heading_level(paragraphs[2]) == SyllabusChunker.WEEK_TITLE_LEVEL

    assert chunker._heading_level(_FakeParagraph("x", "标题 3")) == 3
    assert chunker._heading_level(_FakeParagraph("x", "Heading 4")) == 4


def test_heading_stack_pops_same_or_higher_levels(tmp_path: Path) -> None:
    path = _write_paragraphs(
        tmp_path / "stack.docx",
        [
            "一、课程内容",
            "正文A。",
            "（一）教学进度",
            "正文B。",
            "（二）教学方法",
            "正文C。",
            "3.2 索引结构",
            "正文D。",
            "1. 教材",
            "正文E。",
            "第3章 存储与索引",
            "正文F。",
            "3.2 索引结构",
            "正文G。",
            "第4章 查询优化",
            "正文H。",
            "课程目标",
            "正文I。",
            "第一周 绪论",
            "正文J。",
        ],
    )
    chunks = SyllabusChunker().parse_syllabus(path)
    paths = {chunk["content"]: chunk["heading_path"] for chunk in chunks}

    assert paths["正文A。"] == ["一、课程内容"]
    assert paths["正文B。"] == ["一、课程内容", "（一）教学进度"]
    # 同级替换：教学方法取代教学进度，不产生兄弟层级
    assert paths["正文C。"] == ["一、课程内容", "（二）教学方法"]
    assert paths["正文D。"] == ["一、课程内容", "3.2 索引结构"]
    assert paths["正文E。"] == ["一、课程内容", "3.2 索引结构", "1. 教材"]
    # 更高级标题弹出整条子树
    assert paths["正文F。"] == ["第3章 存储与索引"]
    assert paths["正文G。"] == ["第3章 存储与索引", "3.2 索引结构"]
    assert paths["正文H。"] == ["第4章 查询优化"]
    # 关键词标题继承栈顶层级（同级替换）
    assert paths["正文I。"] == ["课程目标"]
    assert paths["正文J。"] == ["课程目标", "第一周 绪论"]


def test_table_sections_inherit_stack_and_append_chapter(tmp_path: Path) -> None:
    path = tmp_path / "schedule.docx"
    document = Document()
    document.add_paragraph("二、课程基本内容")
    document.add_paragraph("正文甲。")
    document.add_paragraph("（一）教学进度")
    table = document.add_table(rows=4, cols=3)
    _fill_table(
        table,
        [
            ["章节次序及名称(细化至节，必填)", "主要教学内容(必填)", "所需学时(必填)"],
            ["第1章 绪论", "内容：基本概念。", "4"],
            ["第2章 模型", "内容：建模方法。", "2"],
            ["复习", "内容：回顾全部章节。", "1"],
        ],
    )
    document.add_paragraph("（二）教学方法")
    document.add_paragraph("正文乙。")
    document.save(path)

    chunks = SyllabusChunker().parse_syllabus(path)
    schedule_chunks = [
        chunk
        for chunk in chunks
        if chunk["section_type"] == "teaching_schedule"
    ]
    assert [chunk["heading_path"] for chunk in schedule_chunks] == [
        ["二、课程基本内容", "（一）教学进度", "第1章 绪论"],
        ["二、课程基本内容", "（一）教学进度", "第2章 模型"],
        ["二、课程基本内容", "（一）教学进度", "复习"],
    ]
    assert [chunk["chapter_no"] for chunk in schedule_chunks] == [1, 2, None]
    # M1.4：表头括注被剥离，标签换成短名（供 doc_tree 解析学时）
    assert "学时：4" in schedule_chunks[0]["content"]
    assert "章节：第1章 绪论" in schedule_chunks[0]["content"]
    assert "(必填)" not in "\n".join(chunk["content"] for chunk in chunks)

    # 表格块不入栈：后面的（二）教学方法仍是「二、课程基本内容」的子节点
    methods = next(chunk for chunk in chunks if chunk["content"] == "正文乙。")
    assert methods["heading_path"] == ["二、课程基本内容", "（二）教学方法"]


def test_table_without_heading_falls_back_to_own_title(tmp_path: Path) -> None:
    path = tmp_path / "table-only.docx"
    document = Document()
    table = document.add_table(rows=2, cols=4)
    _fill_table(
        table,
        [
            ["课程代码", "IM999", "课程名称", "测试课程"],
            ["课程负责人", "张三", "学分", "3"],
        ],
    )
    document.save(path)

    chunks = SyllabusChunker().parse_syllabus(path)
    assert chunks
    assert all(chunk["heading_path"] == ["课程基本信息"] for chunk in chunks)


def test_build_syllabus_chunks_freezes_structure_fields(tmp_path: Path) -> None:
    syllabus_root = tmp_path / "syllabi"
    syllabus_root.mkdir()
    path = syllabus_root / "IM999测试课程.docx"
    document = Document()
    document.add_paragraph("测试课程教学大纲")
    document.add_paragraph("一、课程基本信息")
    basic_table = document.add_table(rows=3, cols=4)
    _fill_table(
        basic_table,
        [
            ["课程代码", "IM999", "课程名称", "测试课程"],
            ["课程负责人", "张三", "学分", "3"],
            ["先修课程", "信息管理学基础", "学时", "54"],
        ],
    )
    document.add_paragraph("二、课程内容")
    document.add_paragraph("正文甲。")
    document.add_paragraph("（一）教学进度")
    schedule_table = document.add_table(rows=2, cols=3)
    _fill_table(
        schedule_table,
        [
            ["章节次序及名称(细化至节，必填)", "主要教学内容(必填)", "所需学时(必填)"],
            ["第1章 绪论", "内容：基本概念。", "4"],
        ],
    )
    document.add_paragraph("（二）教学方法")
    document.add_paragraph("本课程采用课堂讲授与案例研讨相结合的方式。" * 60)
    document.save(path)

    catalog = {
        "IM999": {
            "course_code": "IM999",
            "course_name": "测试课程",
            "credits": 3,
            "prerequisites": [],
        }
    }
    chunks = SyllabusChunker().build_syllabus_chunks(path, syllabus_root, catalog)

    assert chunks
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(len(chunks)))
    frozen_keys = (
        "chunk_index",
        "heading_path",
        "breadcrumb",
        "parent_id",
        "chapter_no",
        "is_header",
        "prev_chunk_id",
        "next_chunk_id",
    )
    for chunk in chunks:
        for key in frozen_keys:
            assert key in chunk, key
            assert key in chunk["metadata"], key
        assert chunk["metadata"]["heading_path"] == chunk["heading_path"]
        assert chunk["metadata"]["breadcrumb"] == chunk["breadcrumb"]
        assert chunk["metadata"]["chunk_index"] == chunk["chunk_index"]
        assert len(chunk["text"]) <= SyllabusChunker.MAX_TEXT_CHARS
        assert chunk["breadcrumb"].startswith("【测试课程 | ")
        assert chunk["breadcrumb"].endswith("】")

    # prev/next 按 chunk_index 相邻，首尾为 None 且与 metadata 同步
    assert chunks[0]["prev_chunk_id"] is None
    assert chunks[-1]["next_chunk_id"] is None
    for previous, following in zip(chunks, chunks[1:]):
        assert previous["next_chunk_id"] == following["chunk_id"]
        assert following["prev_chunk_id"] == previous["chunk_id"]
        assert following["metadata"]["prev_chunk_id"] == previous["chunk_id"]

    # 头部块：文档概览 + basic_info
    assert chunks[0]["heading_path"] == ["文档概览"]
    assert chunks[0]["is_header"] is True
    assert chunks[0]["parent_id"] is None
    assert chunks[0]["breadcrumb"] == "【测试课程 | 文档概览】"
    basic_info = next(
        chunk
        for chunk in chunks
        if chunk["metadata"]["section_type"] == "basic_info"
    )
    assert basic_info["is_header"] is True
    assert basic_info["heading_path"] == ["一、课程基本信息"]

    document_hash = chunks[0]["document_hash"]
    schedule = next(chunk for chunk in chunks if chunk["chapter_no"] == 1)
    assert schedule["heading_path"] == ["二、课程内容", "（一）教学进度", "第1章 绪论"]
    assert schedule["breadcrumb"] == "【测试课程 | （一）教学进度 | 第1章 绪论】"
    assert schedule["text"].startswith("（一）教学进度 / 第1章 绪论\n\n")
    assert schedule["parent_id"] == node_id(
        document_hash,
        ["二、课程内容", "（一）教学进度"],
    )
    assert schedule["metadata"]["parent_id"] == schedule["parent_id"]
    assert schedule["chapter_no"] == 1
    assert schedule["metadata"]["chapter_no"] == 1

    # 多块 split 出来的部分逐个递增 chunk_index，且标题前缀是最近两级路径
    methods = [chunk for chunk in chunks if chunk["section"] == "教学方法"]
    assert len(methods) >= 2
    assert methods[0]["heading_path"] == ["二、课程内容", "（二）教学方法"]
    assert all(
        chunk["text"].startswith("二、课程内容 / （二）教学方法\n\n")
        for chunk in methods
    )
    assert [chunk["chunk_index"] for chunk in methods] == list(
        range(methods[0]["chunk_index"], methods[0]["chunk_index"] + len(methods))
    )
