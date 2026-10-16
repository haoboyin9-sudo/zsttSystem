"""大纲结构化字段（M3.1 / M3.2）的规则抽取测试。

fixture 用 python-docx 现场生成或直接构造 chunk 字典，不依赖 ``data/`` 语料；
覆盖点来自改造方案 §7 ``tests/test_basic_info_profile.py`` 一行：键名归一、
多教师拆分、``/`` 视为缺失。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from src.data_processing.parser_chunker import (
    SyllabusChunker,
    collect_syllabus_profile,
    parse_basic_info,
)

PROFILE_FIELDS = {
    "course_owner",
    "lecturers",
    "credits",
    "hours",
    "theory_hours",
    "practice_hours",
    "grade",
    "audience",
    "department",
    "intended_semester",
    "prerequisites_raw",
}


def test_parse_basic_info_normalizes_keys_and_splits_lecturers() -> None:
    text = "\n".join(
        [
            "一、课程基本信息",
            "课程负责人：张三",
            "主讲教师：李四，王五、赵六",
            "学分：3",
            "总学时：54",
            "理论学时：36",
            "实践（含实验）学时：18",
            "授课年级：大一",
            "面向专业/大类：图书情报与档案管理类",
            "开课单位：信息管理学院",
            "意向开课学期：2025-1",
            "先修课程：《信息管理学基础》",
        ]
    )

    profile = parse_basic_info(text)

    assert set(profile) == PROFILE_FIELDS
    assert profile["course_owner"] == "张三"
    assert profile["lecturers"] == ["李四", "王五", "赵六"]
    assert profile["credits"] == "3"
    assert profile["hours"] == "54"
    assert profile["theory_hours"] == "36"
    assert profile["practice_hours"] == "18"
    assert profile["grade"] == "大一"
    assert profile["audience"] == "图书情报与档案管理类"
    assert profile["department"] == "信息管理学院"
    assert profile["intended_semester"] == "2025-1"
    assert profile["prerequisites_raw"] == "《信息管理学基础》"


def test_parse_basic_info_accepts_key_aliases() -> None:
    profile = parse_basic_info(
        "\n".join(
            [
                "任课教师：甲、乙",
                "学时：48",
                "适用专业：信息管理与信息系统",
            ]
        )
    )

    assert profile["lecturers"] == ["甲", "乙"]
    assert profile["hours"] == "48"
    assert profile["audience"] == "信息管理与信息系统"
    assert profile["course_owner"] is None


def test_parse_basic_info_treats_placeholders_as_missing() -> None:
    profile = parse_basic_info(
        "\n".join(
            [
                "课程负责人：/",
                "主讲教师：无",
                "学分：-",
                "总学时：—",
                "先修课程：暂无",
                "授课年级：　",
                "开课单位：",
            ]
        )
    )

    for key in (
        "course_owner",
        "lecturers",
        "credits",
        "hours",
        "prerequisites_raw",
        "grade",
        "department",
    ):
        assert profile[key] is None, key
    assert profile["audience"] is None


def test_parse_basic_info_keeps_first_non_missing_value() -> None:
    profile = parse_basic_info("课程负责人：/\n课程负责人：张三\n学分：3\n学分：4")

    assert profile["course_owner"] == "张三"
    assert profile["credits"] == "3"


def _chunk(chunk_id: str, section_type: str, text: str) -> dict[str, Any]:
    return {
        "chunk_id": chunk_id,
        "text": text,
        "metadata": {"section_type": section_type},
    }


def test_collect_syllabus_profile_reads_basic_info_and_assessment_chunks() -> None:
    chunks = [
        _chunk(
            "c1",
            "basic_info",
            "一、课程基本信息\n\n课程负责人：张三\n主讲教师：李四、王五\n"
            "学分：3\n先修课程：/",
        ),
        _chunk(
            "c2",
            "basic_info",
            "一、课程基本信息\n\n学时：54\n理论学时：36\n实践（含实验）学时：18",
        ),
        _chunk(
            "c3",
            "assessment",
            "六、成绩评定方式\n\n考 核 方 式：考试（闭卷）\n成绩录入方式：百分制",
        ),
        _chunk("c4", "teaching_content", "三、教学内容\n\n第一章 绪论"),
    ]

    profile = collect_syllabus_profile(
        chunks,
        source_file="2025-1教学大纲汇总/IM999测试课程.docx",
        course_code="IM999",
        course_name="测试课程",
        source_year="2025",
    )

    assert profile["course_owner"] == "张三"
    assert profile["lecturers"] == ["李四", "王五"]
    assert profile["credits"] == "3"
    assert profile["hours"] == "54"
    assert profile["theory_hours"] == "36"
    assert profile["practice_hours"] == "18"
    assert profile["prerequisites_raw"] is None
    assert profile["assessment_mode"] == "考试（闭卷）"
    assert profile["grading_mode"] == "百分制"
    assert profile["source_file"] == "2025-1教学大纲汇总/IM999测试课程.docx"
    assert profile["course_code"] == "IM999"
    assert profile["course_name"] == "测试课程"
    assert profile["source_year"] == "2025"
    assert profile["field_chunk_ids"] == {
        "course_owner": "c1",
        "lecturers": "c1",
        "credits": "c1",
        "hours": "c2",
        "theory_hours": "c2",
        "practice_hours": "c2",
        "assessment_mode": "c3",
        "grading_mode": "c3",
    }


def test_collect_syllabus_profile_is_empty_without_matching_chunks() -> None:
    profile = collect_syllabus_profile(
        [_chunk("c1", "teaching_content", "三、教学内容\n\n第一章 绪论")],
        source_file="a.docx",
        course_code="IM000",
        course_name="",
        source_year="",
    )

    assert all(profile[key] is None for key in PROFILE_FIELDS)
    assert profile["assessment_mode"] is None
    assert profile["grading_mode"] is None
    assert profile["field_chunk_ids"] == {}


def test_syllabus_profile_end_to_end_from_docx(tmp_path: Path) -> None:
    syllabus_root = tmp_path / "syllabi"
    syllabus_root.mkdir()
    path = syllabus_root / "IM999测试课程.docx"
    document = Document()
    table = document.add_table(rows=3, cols=4)
    rows = [
        ["课程代码", "IM999", "课程名称", "测试课程"],
        ["课程负责人", "张三", "学分", "3"],
        ["主讲教师", "李四、王五", "学时", "54"],
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("三、成绩评定方式")
    document.add_paragraph("考 核 方 式：考试（闭卷）")
    document.add_paragraph("成绩录入方式：百分制")
    document.save(path)

    chunks = SyllabusChunker().build_syllabus_chunks(
        path,
        syllabus_root,
        {
            "IM999": {
                "course_code": "IM999",
                "course_name": "测试课程",
                "credits": 3,
                "prerequisites": [],
            }
        },
    )
    profile = collect_syllabus_profile(
        chunks,
        source_file="IM999测试课程.docx",
        course_code="IM999",
        course_name="测试课程",
        source_year="2025",
    )

    assert profile["course_owner"] == "张三"
    assert profile["lecturers"] == ["李四", "王五"]
    assert profile["credits"] == "3"
    assert profile["hours"] == "54"
    assert profile["assessment_mode"] == "考试（闭卷）"
    assert profile["grading_mode"] == "百分制"

    chunk_ids = {chunk["chunk_id"] for chunk in chunks}
    assert set(profile["field_chunk_ids"].values()) <= chunk_ids
    assert profile["field_chunk_ids"]["assessment_mode"] == next(
        chunk["chunk_id"]
        for chunk in chunks
        if chunk["metadata"]["section_type"] == "assessment"
    )
